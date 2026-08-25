#!/usr/bin/env python3
"""수기 회계장부 엑셀의 왼쪽 블록에 `분류` 열(D)을 끼우고 추정 분류를 채운다.

    python3 scripts/add_category_column.py 원본.xlsx -o 결과.xlsx

끼운 뒤 왼쪽 블록은 `read_ledger.py` 가 기대하는 순서가 된다:

    월 · 일 · 적요 · **분류** · 수입 · 지출 · 잔액

## 왜 openpyxl 로 다시 쓰지 않고 XML 을 직접 고치나

openpyxl 로 통째로 다시 쓰면 **수식의 캐시값이 전부 날아간다.** 엑셀에서 열면
다시 계산되지만, 그 전까지 미리보기·Quick Look 에서는 숫자가 빈칸으로 보인다.
장부는 숫자를 보려고 여는 파일이라 그건 손해다. 그래서 셀을 옮기고 참조만
고쳐 쓴다 — 서식·테두리·통화형식은 손대지 않는다.

## 열 삽입이 안전한 이유 (2026-08-21 전수 조사)

수식 438개를 전부 훑어 확인했다:

- 시트 간 참조 **0건** — 다른 시트를 보는 수식이 없다
- 절대참조(`$`) **0건**
- 열을 걸치는 범위(`SUM(D4:F4)` 같은 것) **0건** — 범위는 전부 같은 열 안이다

그래서 **D 이상인 열 참조를 전부 +1 하면 끝난다.** 엑셀이 열을 끼울 때 하는 일과
같다. 이 전제가 깨지면(예: 나중에 시트 간 참조가 생기면) 이 스크립트를 다시 쓰기
전에 조사부터 해야 한다 — `--audit` 으로 다시 확인할 수 있다.

## 건드리는 시트

3행 머리글이 `월·일·적요·수입·지출·잔액` 인 시트만. `결산내역서`(레이아웃이 다르다)
와 `그라운드 룰`(수식 없음)은 건드리지 않는다.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl 이 필요하다:  python3 -m pip install openpyxl")

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from classify_ledger import classify  # noqa: E402

INSERT_AT = 4  # D열
LEFT_HEADER = ["월", "일", "적요", "수입", "지출", "잔액"]
NEW_HEADER = "분류"
NEW_WIDTH = "14"


def col_to_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def num_to_col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def shift(letters: str) -> str:
    """D 이상이면 한 칸 오른쪽."""
    n = col_to_num(letters)
    return num_to_col(n + 1) if n >= INSERT_AT else letters


def shift_formula(f: str) -> str:
    return re.sub(r"\b([A-Z]{1,3})(\d+)\b", lambda m: shift(m.group(1)) + m.group(2), f)


def shift_range(ref: str) -> str:
    """`A1:J1` 같은 범위. 시작이 삽입점 왼쪽이면 그대로 두고 끝만 민다."""
    def one(part: str) -> str:
        m = re.match(r"([A-Z]+)(\d+)$", part)
        return shift(m.group(1)) + m.group(2)
    return ":".join(one(p) for p in ref.split(":"))


# ── 시트 XML 변환 ──────────────────────────────────────────────────────────

CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
ROW_RE = re.compile(r'(<row [^>]*?r="(\d+)"[^>]*>)(.*?)(</row>)', re.S)
F_RE = re.compile(r"<f([^>]*?)(?:/>|>(.*?)</f>)", re.S)


def shift_formula_tag(inner: str) -> str:
    """`<f>` 하나를 옮긴다.

    **공유 수식(`t="shared"`)이 함정이다.** 마스터만 수식 텍스트를 갖고
    (`<f t="shared" ref="F6:F23" si="0">F5+D6-E6</f>`), 자식은 텍스트 없이
    `si` 로 물려받는다 (`<f t="shared" si="0"/>`). 그래서 **텍스트와 `ref` 범위를
    둘 다 옮겨야 하고**, 자식은 건드리면 안 된다 — 마스터에서 자동으로 따라온다.
    속성 없는 `<f>` 만 고치면 잔액 열이 통째로 옛 열을 가리킨 채 남는다.
    """
    def one(m):
        attrs, text = m.group(1), m.group(2)
        attrs = re.sub(r'ref="([^"]+)"',
                       lambda rm: 'ref="' + shift_range(rm.group(1)) + '"', attrs)
        if text is None:
            return f"<f{attrs}/>"
        return f"<f{attrs}>{shift_formula(text)}</f>"
    return F_RE.sub(one, inner)


def transform_sheet(xml: str, categories: dict[int, str]) -> str:
    def do_row(m):
        open_tag, row_no, body, close_tag = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        cells = []
        c_style = ""
        for cm in CELL_RE.finditer(body):
            letters, rno, attrs, inner = cm.group(1), cm.group(2), cm.group(3), cm.group(4)
            # 새 D 셀 서식은 같은 줄 적요(C)에서 물려받는다 — XML 의 s= 를 그대로 쓴다
            if letters == "C":
                sm = re.search(r's="(\d+)"', attrs)
                c_style = sm.group(1) if sm else ""
            new_col = shift(letters)
            if inner is not None:
                cell = f'<c r="{new_col}{rno}"{attrs}>{shift_formula_tag(inner)}</c>'
            else:
                cell = f'<c r="{new_col}{rno}"{attrs}/>'
            cells.append((col_to_num(new_col), cell))

        # 새 D 셀. 서식은 같은 줄 적요(C)에서 물려받아 테두리가 이어지게 한다.
        value = categories.get(row_no)
        style_attr = f' s="{c_style}"' if c_style else ""
        if value:
            escaped = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            new_cell = f'<c r="D{row_no}"{style_attr} t="inlineStr"><is><t>{escaped}</t></is></c>'
        else:
            new_cell = f'<c r="D{row_no}"{style_attr}/>'
        cells.append((INSERT_AT, new_cell))

        cells.sort(key=lambda x: x[0])
        open_tag = re.sub(r'spans="1:(\d+)"',
                          lambda sm: f'spans="1:{int(sm.group(1)) + 1}"', open_tag)
        return open_tag + "".join(c for _, c in cells) + close_tag

    xml = ROW_RE.sub(do_row, xml)

    # 시트 범위
    xml = re.sub(r'(<dimension ref=")([^"]+)(")',
                 lambda m: m.group(1) + shift_range(m.group(2)) + m.group(3), xml)

    # 병합 셀
    xml = re.sub(r'(<mergeCell ref=")([^"]+)(")',
                 lambda m: m.group(1) + shift_range(m.group(2)) + m.group(3), xml)

    # 열 너비: D 이상을 밀고, 새 D 너비를 넣는다
    def do_cols(m):
        body = m.group(2)
        def one(cm):
            mn, mx = int(cm.group(1)), int(cm.group(2))
            nmn = mn + 1 if mn >= INSERT_AT else mn
            nmx = mx + 1 if mx >= INSERT_AT else mx
            return cm.group(0).replace(f'min="{mn}"', f'min="{nmn}"').replace(f'max="{mx}"', f'max="{nmx}"')
        body = re.sub(r'<col min="(\d+)" max="(\d+)"[^>]*/>', one, body)
        body = f'<col min="{INSERT_AT}" max="{INSERT_AT}" width="{NEW_WIDTH}" customWidth="1"/>' + body
        return m.group(1) + body + m.group(3)

    xml = re.sub(r"(<cols>)(.*?)(</cols>)", do_cols, xml, flags=re.S)
    return xml


# ── 준비: 어느 시트를, 어떤 분류로 ─────────────────────────────────────────

def plan(path: str):
    """{시트이름: {행: 분류}} 를 만든다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    targets = {}
    for ws in wb.worksheets:
        header = [ws.cell(3, i).value for i in range(1, 7)]
        if header != LEFT_HEADER:
            continue
        cats: dict[int, str] = {3: NEW_HEADER}
        for row in range(1, ws.max_row + 1):
            desc = ws.cell(row, 3).value
            if row < 4 or not desc:
                continue
            desc = str(desc).strip()
            if desc == "전월이월":
                continue
            income, expense = ws.cell(row, 4).value, ws.cell(row, 5).value
            has_i = isinstance(income, (int, float)) and income
            has_e = isinstance(expense, (int, float)) and expense
            if has_i and has_e:
                sys.exit(f"{ws.title} {row}행: 수입과 지출이 같이 있다 — 손으로 봐야 한다")
            if has_i or has_e:
                cats[row] = classify(desc, bool(has_i))
        targets[ws.title] = cats
    return targets


def sheet_paths(parts: dict) -> dict[str, str]:
    rels = parts["xl/_rels/workbook.xml.rels"].decode("utf8")
    rid = {m.group(1): m.group(2) for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    book = parts["xl/workbook.xml"].decode("utf8")
    return {m.group(1): "xl/" + rid[m.group(2)]
            for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="([^"]+)"', book)}


def audit(path: str) -> list[str]:
    """열 삽입 전제가 아직 성립하는지 확인한다."""
    wb = openpyxl.load_workbook(path)
    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                where = f"{ws.title}!{c.coordinate}"
                if "!" in v:
                    problems.append(f"{where}: 시트 간 참조  {v}")
                if "$" in v:
                    problems.append(f"{where}: 절대참조  {v}")
                for a, _, b, _ in re.findall(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", v):
                    if a != b:
                        problems.append(f"{where}: 열을 걸치는 범위  {v}")
    return problems


def cells_read_by(formula: str) -> list[str]:
    """수식이 읽는 셀 좌표들. 범위는 펼친다."""
    out = []
    for m in re.finditer(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", formula):
        col, r1, r2 = m.group(1), int(m.group(2)), int(m.group(4))
        out += [f"{col}{r}" for r in range(r1, r2 + 1)]
    rest = re.sub(r"[A-Z]+\d+:[A-Z]+\d+", "", formula)
    out += [m.group(0) for m in re.finditer(r"\b[A-Z]{1,3}\d+\b", rest)]
    return out


def verify(before: str, after: str) -> list[str]:
    """옮긴 결과가 옳은지 원본과 대조한다.

    수식은 **텍스트가 아니라 읽는 값으로** 비교한다. 참조가 바뀌는 게 이 작업의
    목적이라 텍스트 비교는 아무것도 증명하지 못한다 — 참조를 풀어서 같은 데이터를
    읽는지 봐야 공유 수식을 빠뜨린 것 같은 사고가 잡힌다.
    """
    bf, bv = openpyxl.load_workbook(before), openpyxl.load_workbook(before, data_only=True)
    af_, av = openpyxl.load_workbook(after), openpyxl.load_workbook(after, data_only=True)
    problems = []
    if bf.sheetnames != af_.sheetnames:
        return ["시트 목록이 달라졌다"]

    for name in bf.sheetnames:
        ob, obv, nb, nbv = bf[name], bv[name], af_[name], av[name]
        moved = [obv.cell(3, i).value for i in range(1, 7)] == LEFT_HEADER
        for row in range(1, ob.max_row + 1):
            for col in range(1, 11):
                new_col = col + 1 if (moved and col >= INSERT_AT) else col
                old, new = obv.cell(row, col).value, nbv.cell(row, new_col).value
                if old != new:
                    problems.append(
                        f"{name}!{obv.cell(row, col).coordinate}: 값이 달라졌다 {old!r} → {new!r}")
        for r in ob.iter_rows():
            for c in r:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                new_col = c.column + 1 if (moved and c.column >= INSERT_AT) else c.column
                nc = nb.cell(c.row, new_col)
                want = [obv[x].value for x in cells_read_by(c.value)]
                got = [nbv[x].value for x in cells_read_by(nc.value)]
                if want != got:
                    problems.append(
                        f"{name}!{c.coordinate}: 수식이 읽는 값이 달라졌다  "
                        f"{c.value} → {nc.value}")
    return problems


def main() -> None:
    ap = argparse.ArgumentParser(description="회계장부에 분류 열(D)을 끼운다")
    ap.add_argument("path")
    ap.add_argument("-o", "--out", help="결과 파일 (없으면 원본 옆에 ' - 분류' 를 붙인다)")
    ap.add_argument("--audit", action="store_true", help="열 삽입 전제만 확인하고 끝낸다")
    args = ap.parse_args()

    problems = audit(args.path)
    if problems:
        print("열 삽입 전제가 깨졌다. 아래를 먼저 해결할 것:", file=sys.stderr)
        for p in problems:
            print("  •", p, file=sys.stderr)
        sys.exit(1)
    print(f"전수 조사 통과 — 시트 간 참조·절대참조·열을 걸치는 범위 모두 없음")
    if args.audit:
        return

    out = args.out or args.path.rsplit(".", 1)[0] + " - 분류.xlsx"
    targets = plan(args.path)

    zin = zipfile.ZipFile(args.path)
    parts = {n: zin.read(n) for n in zin.namelist()}
    order = zin.namelist()
    zin.close()
    paths = sheet_paths(parts)

    filled = defaultdict(int)
    for name, cats in targets.items():
        p = paths[name]
        parts[p] = transform_sheet(parts[p].decode("utf8"), cats).encode("utf8")
        filled[name] = sum(1 for r, v in cats.items() if r != 3 and v)

    # 수식의 참조가 바뀌었으니 계산 순서 캐시는 버린다. 엑셀이 다시 만든다.
    drop = "xl/calcChain.xml"
    if drop in parts:
        del parts[drop]
        order = [n for n in order if n != drop]
        ct = parts["[Content_Types].xml"].decode("utf8")
        parts["[Content_Types].xml"] = re.sub(
            r'<Override[^>]*calcChain[^>]*/>', "", ct).encode("utf8")
        rels = parts["xl/_rels/workbook.xml.rels"].decode("utf8")
        parts["xl/_rels/workbook.xml.rels"] = re.sub(
            r'<Relationship[^>]*calcChain[^>]*/>', "", rels).encode("utf8")

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in order:
            z.writestr(n, parts[n])

    print(f"\n분류 열을 끼운 시트 {len(targets)}개:")
    for name in sorted(targets):
        print(f"  {name:<10} 분류 채운 줄 {filled[name]:>3}")

    problems = verify(args.path, out)
    if problems:
        os.remove(out)
        print(f"\n검증 실패 {len(problems)}건 — 결과 파일을 지웠다.", file=sys.stderr)
        for p in problems[:20]:
            print("  •", p, file=sys.stderr)
        if len(problems) > 20:
            print(f"  … 외 {len(problems) - 20}건", file=sys.stderr)
        sys.exit(1)

    print("검증 통과 — 모든 값이 제자리로 옮겨졌고, 모든 수식이 같은 값을 읽는다")
    print(f"\n저장: {out}")


if __name__ == "__main__":
    main()
